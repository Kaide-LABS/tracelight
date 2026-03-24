import os
import json
from docxtpl import DocxTemplate
from pptx import Presentation
from app.schemas_v2 import MemoSection, CitedMetrics

class CitationFormatter:
    def __init__(self, templates_dir: str, output_dir: str):
        self.templates_dir = templates_dir
        self.output_dir = output_dir

    def format_deliverables(self, job_id: str, sections: list[MemoSection], metrics: CitedMetrics):
        os.makedirs(self.output_dir, exist_ok=True)
        docx_path = os.path.join(self.output_dir, f"{job_id}.docx")
        pptx_path = os.path.join(self.output_dir, f"{job_id}.pptx")
        json_path = os.path.join(self.output_dir, f"{job_id}_audit.json")

        # 1. DOCX
        tpl = DocxTemplate(os.path.join(self.templates_dir, "ic_memo_template.docx"))
        context = {sec.section_id: sec.content for sec in sections}
        tpl.render(context)
        tpl.save(docx_path)

        # 2. PPTX
        prs = Presentation(os.path.join(self.templates_dir, "exec_deck_template.pptx"))
        for sec in sections:
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            body = slide.shapes.placeholders[1]
            title.text = sec.title
            body.text = sec.content[:500] + ("..." if len(sec.content) > 500 else "")
        prs.save(pptx_path)

        # 3. Audit Trail JSON
        audit_trail = []
        for sec in sections:
            for cit in sec.citations:
                audit_trail.append({
                    "section": sec.section_id,
                    "citation": cit,
                    "source": "See CitedMetrics or ContextHarvester for details."
                })
        
        with open(json_path, 'w') as f:
            json.dump(audit_trail, f, indent=2)

        return {
            "docx": f"/api/v2/memo/download/{job_id}?fmt=docx",
            "pptx": f"/api/v2/memo/download/{job_id}?fmt=pptx",
            "audit_trail": f"/api/v2/memo/download/{job_id}?fmt=json"
        }
