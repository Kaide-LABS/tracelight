import csv
import json
import io
import os
from typing import List, Dict
import docx
from app.schemas_v3 import DraftResponse, ComplianceConfig

class RoutingExporterAgent:
    def __init__(self, output_dir: str):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def route_responses(self, responses: List[DraftResponse], config: ComplianceConfig) -> List[DraftResponse]:
        routed = []
        for r in responses:
            if r.confidence >= config.auto_approve_above:
                r.status = "auto_approved"
            elif r.confidence < config.confidence_threshold:
                r.status = "needs_review"
            else:
                r.status = "needs_review"  # Conservative routing
            routed.append(r)
        return routed

    def export_csv(self, job_id: str, responses: List[DraftResponse]) -> str:
        filepath = os.path.join(self.output_dir, f"{job_id}_questionnaire.csv")
        with open(filepath, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(["question_id", "domain", "question_text", "response", "confidence", "status", "citations"])
            for r in responses:
                writer.writerow([
                    r.question_id,
                    r.domain,
                    r.question_text,
                    r.response_text,
                    f"{r.confidence:.2f}",
                    r.status,
                    ", ".join(r.citations)
                ])
        return filepath

    def export_json(self, job_id: str, responses: List[DraftResponse]) -> str:
        filepath = os.path.join(self.output_dir, f"{job_id}_questionnaire.json")
        data = [r.model_dump() for r in responses]
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2)
        return filepath

    def export_docx(self, job_id: str, responses: List[DraftResponse]) -> str:
        filepath = os.path.join(self.output_dir, f"{job_id}_questionnaire.docx")
        doc = docx.Document()
        doc.add_heading('Compliance Questionnaire Responses', 0)
        
        # Group by domain
        domains: Dict[str, List[DraftResponse]] = {}
        for r in responses:
            domain = r.domain or "Uncategorized"
            if domain not in domains:
                domains[domain] = []
            domains[domain].append(r)
            
        for domain, items in domains.items():
            doc.add_heading(domain, level=1)
            for r in items:
                p_q = doc.add_paragraph()
                p_q.add_run(f"[{r.question_id}] ").bold = True
                p_q.add_run(r.question_text)
                
                p_a = doc.add_paragraph()
                p_a.add_run("Response:\n").bold = True
                p_a.add_run(r.response_text)
                
                p_meta = doc.add_paragraph()
                p_meta.add_run(f"Confidence: {r.confidence:.2f} | Status: {r.status}").italic = True
                if r.citations:
                    p_meta.add_run(f"\nCitations: {', '.join(r.citations)}")
                    
                doc.add_paragraph("-" * 40)
                
        doc.save(filepath)
        return filepath
