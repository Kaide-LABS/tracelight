import os
import uuid
import httpx
from fastapi import FastAPI, Depends, UploadFile, File, BackgroundTasks, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from slowapi.errors import RateLimitExceeded
from slowapi import _rate_limit_exceeded_handler
from app.schemas import GenerateRequest, GenerateResponse
from app.config import Settings
from app.pipeline import run_pipeline
from app.templates import PRESETS
from app.auth import verify_api_key
from app.database import init_db, get_job
from app.logging_config import setup_logging, get_logger
from app.middleware import limiter, RequestLoggingMiddleware

# Phase 2 imports
from app.schemas_v2 import MemoGenerateRequest, MemoGenerateResponse
from app.pipeline_v2 import run_memo_pipeline, get_chroma_client
from app.agents.context_harvester import ContextHarvester

app = FastAPI(title="Tracelight Synthetic Data Generator API")

app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)
app.add_middleware(RequestLoggingMiddleware)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def get_settings():
    return Settings()

log = get_logger("error_handler")

@app.on_event("startup")
async def startup_event():
    settings = get_settings()
    os.makedirs(settings.output_dir, exist_ok=True)
    os.makedirs("app/templates_v2", exist_ok=True)
    
    # Initialize SQLite
    await init_db()

    # Initialize logging
    setup_logging(settings.log_level)
    
    # Generate default templates if missing
    import docx
    from pptx import Presentation
    
    docx_path = "app/templates_v2/ic_memo_template.docx"
    if not os.path.exists(docx_path):
        doc = docx.Document()
        doc.add_heading('Investment Committee Memorandum', 0)
        for sec in ['executive_summary', 'investment_thesis', 'market_analysis', 'financial_projections', 'deal_structure', 'risk_mitigation']:
            doc.add_heading(sec.replace("_", " ").title(), level=1)
            doc.add_paragraph(f"{{{{ {sec} }}}}")
        doc.save(docx_path)
        
    pptx_path = "app/templates_v2/exec_deck_template.pptx"
    if not os.path.exists(pptx_path):
        prs = Presentation()
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "Executive Deck Template"
        subtitle.text = "Tracelight Generated"
        prs.save(pptx_path)

@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    log.error(
        "unhandled_exception",
        path=request.url.path,
        method=request.method,
        error=str(exc),
        exc_info=True,
    )
    return JSONResponse(
        status_code=500,
        content={"detail": "Internal server error", "error_id": str(uuid.uuid4())},
    )

@app.post("/api/v1/generate", response_model=GenerateResponse, dependencies=[Depends(verify_api_key)])
@limiter.limit("10/minute")
async def generate_data(request: Request, body: GenerateRequest, settings: Settings = Depends(get_settings)):
    return await run_pipeline(body, settings)

@app.get("/api/v1/download/{job_id}", dependencies=[Depends(verify_api_key)])
@limiter.limit("60/minute")
async def download_data(request: Request, job_id: str, settings: Settings = Depends(get_settings)):
    csv_path = f"{settings.output_dir}/{job_id}.csv"
    json_path = f"{settings.output_dir}/{job_id}.json"
    
    if os.path.exists(csv_path):
        return FileResponse(csv_path, media_type="text/csv", filename=f"{job_id}.csv")
    elif os.path.exists(json_path):
        return FileResponse(json_path, media_type="application/json", filename=f"{job_id}.json")
    
    raise HTTPException(status_code=404, detail="File not found")

@app.get("/api/v1/templates", dependencies=[Depends(verify_api_key)])
@limiter.limit("60/minute")
async def get_templates(request: Request):
    return PRESETS

# ======================= PHASE 2 ROUTES =======================

@app.post("/api/v2/memo/upload-sources", dependencies=[Depends(verify_api_key)])
@limiter.limit("20/minute")
async def upload_sources(request: Request, files: list[UploadFile] = File(...)):
    session_id = str(uuid.uuid4())
    harvester = ContextHarvester(get_chroma_client())
    
    temp_dir = f"/tmp/{session_id}"
    os.makedirs(temp_dir, exist_ok=True)
    
    try:
        for file in files:
            filepath = os.path.join(temp_dir, file.filename)
            with open(filepath, "wb") as f:
                f.write(await file.read())
            harvester.process_file(filepath, session_id)
    finally:
        # Cleanup
        import shutil
        shutil.rmtree(temp_dir, ignore_errors=True)
        
    return {"session_id": session_id}

@app.post("/api/v2/memo/generate", response_model=MemoGenerateResponse, dependencies=[Depends(verify_api_key)])
@limiter.limit("5/minute")
async def generate_memo(request: Request, body: MemoGenerateRequest, background_tasks: BackgroundTasks, settings: Settings = Depends(get_settings)):
    job_id = str(uuid.uuid4())
    
    from app.database import create_job
    await create_job(job_id, "v2", body.model_dump())
    
    # Run async pipeline
    background_tasks.add_task(run_memo_pipeline, job_id, body, settings)
    
    return MemoGenerateResponse(
        job_id=job_id,
        status="processing",
        generated_at=str(uuid.uuid1())
    )

@app.get("/api/v2/memo/status/{job_id}", response_model=MemoGenerateResponse, dependencies=[Depends(verify_api_key)])
@limiter.limit("60/minute")
async def get_memo_status(request: Request, job_id: str):
    job = await get_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
        
    status = job["status"]
    sections = job["result"] if status == "completed" and job.get("result") else None
    
    # We still need URLs if any are generated - pipeline_v2 must store them in result
    urls = job.get("result", {}).get("urls", {}) if isinstance(job.get("result"), dict) else {}
    
    return MemoGenerateResponse(
        job_id=job_id,
        status=status,
        sections=sections.get("sections") if isinstance(sections, dict) and "sections" in sections else (sections if isinstance(sections, list) else None),
        download_urls=urls,
        audit_trail_url=urls.get("audit_trail") if urls else None,
        generated_at=job.get("updated_at", "")
    )

@app.get("/api/v2/memo/download/{job_id}", dependencies=[Depends(verify_api_key)])
@limiter.limit("60/minute")
async def download_memo(request: Request, job_id: str, fmt: str, settings: Settings = Depends(get_settings)):
    if fmt == "docx":
        path = f"{settings.output_dir}/{job_id}.docx"
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif fmt == "pptx":
        path = f"{settings.output_dir}/{job_id}.pptx"
        media = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif fmt == "json":
        path = f"{settings.output_dir}/{job_id}_audit.json"
        media = "application/json"
    else:
        raise HTTPException(status_code=400, detail="Invalid format")
        
    if os.path.exists(path):
        return FileResponse(path, media_type=media, filename=f"{job_id}.{fmt}")
    
    raise HTTPException(status_code=404, detail="File not found")
@app.get("/health")
async def health_check(settings: Settings = Depends(get_settings)):
    checks = {"status": "ok", "checks": {}}

    # SQLite
    try:
        from app.database import get_job
        job = await get_job("__probe__")
        checks["checks"]["database"] = "ok"
    except Exception as e:
        checks["checks"]["database"] = f"error: {e}"
        checks["status"] = "degraded"

    # LLM reachability (only if not demo mode)
    if not settings.demo_mode and settings.llm_api_key:
        try:
            if settings.llm_provider == "google":
                from google import genai
                client = genai.Client(api_key=settings.llm_api_key)
                client.models.list()
                checks["checks"]["llm"] = "ok"
            else:
                import httpx
                resp = httpx.get(
                    f"{settings.llm_base_url}/models",
                    headers={"Authorization": f"Bearer {settings.llm_api_key}"},
                    timeout=5.0,
                )
                checks["checks"]["llm"] = "ok" if resp.status_code == 200 else f"error: {resp.status_code}"
        except Exception as e:
            checks["checks"]["llm"] = f"error: {e}"
            checks["status"] = "degraded"
    else:
        checks["checks"]["llm"] = "skipped (demo mode)"

    # Disk space
    import shutil
    usage = shutil.disk_usage(settings.output_dir)
    free_gb = usage.free / (1024**3)
    checks["checks"]["disk_free_gb"] = round(free_gb, 1)
    if free_gb < 1.0:
        checks["status"] = "degraded"

    return checks
