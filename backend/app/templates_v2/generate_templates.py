import os
import docx
from pptx import Presentation

def generate_templates():
    os.makedirs(os.path.dirname(__file__), exist_ok=True)
    
    # Generate DOCX
    docx_path = os.path.join(os.path.dirname(__file__), "ic_memo_template.docx")
    doc = docx.Document()
    doc.add_heading('Investment Committee Memorandum', 0)
    for sec in ['executive_summary', 'investment_thesis', 'market_analysis', 'financial_projections', 'deal_structure', 'risk_mitigation']:
        doc.add_heading(sec.replace("_", " ").title(), level=1)
        doc.add_paragraph(f"{{{{ {sec} }}}}")
    doc.save(docx_path)
    print(f"Generated {docx_path}")
    
    # Generate PPTX
    pptx_path = os.path.join(os.path.dirname(__file__), "exec_deck_template.pptx")
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Executive Deck Template"
    subtitle.text = "Tracelight Generated"
    prs.save(pptx_path)
    print(f"Generated {pptx_path}")

if __name__ == "__main__":
    generate_templates()
